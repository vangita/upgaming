"""Writes the sample documents in data/samples/.

One fictional incident (INC-2041 at the fictional company Corvane Pay) described by four
teams in four formats. The sources disagree on purpose, to test contradiction detection:

- root cause: expired TLS certificate (post-mortem, Q1 report) vs database failover (status page)
- duration: 47 minutes (post-mortem, Q1 report) vs about 65 minutes (status page)
- failed card transactions: 12% (post-mortem) vs 18% (Q1 report)

No document gives the financial cost of the incident or the name of the on-call engineer,
so questions about them must end in "not enough information".
"""
import logging
from datetime import datetime, timezone
from pathlib import Path

import docx
from fpdf import FPDF
from fpdf.enums import XPos, YPos

from app.core.log import setup_logging

logger = logging.getLogger(__name__)

SAMPLES_DIR = Path(__file__).resolve().parents[1] / "data" / "samples"

POSTMORTEM_MD = """\
# INC-2041 Post-mortem: card payments outage

Author: Site Reliability Engineering team, Corvane Pay. Date: 18 March 2026. Status: final.

## Summary

On 14 March 2026 card payments failed for 47 minutes, from 14:05 to 14:52 UTC. About 12% of all
card transactions in that window were declined with error code GW-503. Bank transfers and wallet
payments were not affected.

## Timeline (UTC)

- 14:05 - The TLS certificate of the internal payment gateway expired. The gateway started
  rejecting connections from the card processing service.
- 14:09 - The alert "card-auth error rate above 5%" fired and paged the on-call engineer.
- 14:31 - The expired certificate was identified as the cause.
- 14:48 - A renewed certificate was deployed to all gateway nodes.
- 14:52 - Error rates returned to normal and the incident was declared resolved.

## Root cause

The certificate had been issued manually in March 2025. The automated renewal job, introduced in
June 2025, only manages certificates created after that date, so this one was never renewed.
No alert existed for certificates close to expiry.

## Action items

1. Move all gateway certificates to automated renewal. Owner: Platform team, due 31 March 2026.
2. Alert 14 days before any certificate expires. Owner: SRE team, due 4 April 2026.
3. Require a status page update within 10 minutes of detection. Owner: Incident management.
"""

STATUS_PAGE_TXT = """\
Corvane Pay status page - incident history export

Incident: Degraded card payments
Affected component: Card payments (EU and US)

14 March 2026, 14:20 UTC - Investigating
We are seeing elevated failure rates for card payments. Our engineers are investigating.

14 March 2026, 14:45 UTC - Identified
The issue has been traced to a database failover in our primary data center. We are working on a fix.

14 March 2026, 15:10 UTC - Resolved
Card payments have been working normally since 15:10 UTC. The total disruption lasted about
65 minutes. We apologise for the inconvenience.
"""

SUPPORT_TITLE = "INC-2041 Customer support summary"
SUPPORT_TICKETS = [("Channel", "Tickets"), ("Email", "610"), ("Chat", "529"), ("Phone", "245")]

REPORT_TITLE = "Reliability report Q1 2026"
REPORT_PAGES = [
    [
        "Corvane Pay - Reliability report Q1 2026",
        "Prepared by Engineering Management on 2 April 2026.",
        "Availability of card payments in Q1 2026 was 99.94%, below the target of 99.95%. "
        "Two incidents caused most of the downtime in the quarter.",
        "INC-2017 (9 February 2026): wallet payments were unavailable for 22 minutes after a faulty "
        "configuration change in the wallet service. The change was rolled back at 09:41 UTC. "
        "Card payments were not affected.",
    ],
    [
        "INC-2041 (14 March 2026): card payments failed for 47 minutes because of an expired TLS "
        "certificate on the internal payment gateway. 18% of card transactions failed during the "
        "incident. The certificate is now covered by automated renewal.",
        "Goals for Q2 2026: alerts for expiring certificates, a status page update within 10 minutes "
        "of detection, and card payment availability of at least 99.95%.",
    ],
]

# Fixed dates keep the document metadata the same on every run
# (the .docx bytes still change: the zip inside stores the time it was written)
SUPPORT_DATE = datetime(2026, 3, 16, 9, 0, tzinfo=timezone.utc)
REPORT_DATE = datetime(2026, 4, 2, 9, 0, tzinfo=timezone.utc)


def write_text(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8", newline="\n")


def write_support_docx(path: Path) -> None:
    document = docx.Document()
    props = document.core_properties
    props.title, props.author = SUPPORT_TITLE, "Customer Support Operations"
    props.created = props.modified = SUPPORT_DATE

    document.add_heading(SUPPORT_TITLE, level=1)
    document.add_paragraph("Prepared by Customer Support Operations on 16 March 2026.")

    document.add_heading("Customer impact", level=2)
    document.add_paragraph(
        "The first customer report about failed card payments arrived at 14:12 UTC on 14 March 2026. "
        "In total 1,384 support tickets were opened about the incident: 912 by merchants and 472 by "
        "cardholders. Tickets by channel:"
    )
    table = document.add_table(rows=len(SUPPORT_TICKETS), cols=2)
    for row, values in zip(table.rows, SUPPORT_TICKETS):
        for cell, value in zip(row.cells, values):
            cell.text = value

    document.add_heading("Merchant compensation", level=2)
    document.add_paragraph(
        "Merchants with more than 50 failed transactions during the incident get a refund of that "
        "day's transaction fees. 214 merchants qualified. Requests must be sent by 30 April 2026."
    )

    document.add_heading("Feedback", level=2)
    document.add_paragraph(
        "Many merchants complained that the status page was first updated 15 minutes after the "
        "failures began."
    )
    document.save(str(path))


def write_report_pdf(path: Path) -> None:
    pdf = FPDF()
    pdf.set_title(REPORT_TITLE)
    pdf.set_author("Engineering Management")
    pdf.set_creation_date(REPORT_DATE)
    pdf.set_font("Helvetica", size=11)
    for paragraphs in REPORT_PAGES:
        pdf.add_page()
        for paragraph in paragraphs:
            pdf.multi_cell(0, 6, paragraph, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.ln(4)
    pdf.output(str(path))


def main() -> None:
    setup_logging()
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)

    write_text(SAMPLES_DIR / "inc-2041-postmortem.md", POSTMORTEM_MD)
    write_text(SAMPLES_DIR / "inc-2041-status-page.txt", STATUS_PAGE_TXT)
    write_support_docx(SAMPLES_DIR / "inc-2041-support-summary.docx")
    write_report_pdf(SAMPLES_DIR / "reliability-report-q1-2026.pdf")
    logger.info("Sample documents written to %s", SAMPLES_DIR)


if __name__ == "__main__":
    main()
